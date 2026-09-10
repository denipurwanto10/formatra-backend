#!/usr/bin/env python3
"""
PDF -> DOCX conversion for the "PDF to Word" tool only.
--------------------------------------------------------
Word -> PDF and PPT -> PDF keep using LibreOffice (server/index.mjs). This
script is invoked ONLY for the /api/convert/pdf-to-word route, as a separate
process, so a problem here (missing Python, bad wheel, a PDF that crashes the
parser) can never touch the other two conversions.

Usage:
    python3 pdf_to_docx.py <input.pdf> <output.docx>

Exit code 0 = success (output.docx was written).
Exit code 1 = failure; a one-line reason is printed to stderr and the Node
server treats that as "engine failed", falling back to the LibreOffice
writer_pdf_import path for this same request automatically.

Why pdf2docx instead of LibreOffice's writer_pdf_import filter:
LibreOffice's PDF import treats a PDF largely as a bag of absolutely-
positioned drawing instructions and re-flows it into ordinary Writer
paragraphs, which is exactly what tends to shift text position, spacing,
alignment, and column layout. pdf2docx instead reconstructs the page as a
layout tree (text blocks with real bounding boxes, table grids, images) and
rebuilds a .docx from that structure — positioned frames/tables and matched
fonts/sizes/spacing instead of a best-effort reflow. That generally keeps the
visual result much closer to the source PDF for resumes, forms, brochures,
multi-column pages, and other precisely laid-out documents.
"""

import logging
import sys


# Bullet glyphs seen from real-world PDFs (common Unicode bullets, plus the
# Wingdings/Symbol private-use codepoints LibreOffice/Word commonly emit for
# list markers, e.g. U+F0B7). Kept in sync with BULLET_RE in the client-side
# pdfToWord.js fallback, so both paths recognize the same set.
_BULLET_CHARS = set("•◦▪‣∙○●■□-*\uf06e\uf0a7\uf0b7\uf0d8\uf0a8")


def _looks_like_bullet_marker(line):
    """True if a line is nothing but a single bullet glyph (with padding)."""
    stripped = line.strip()
    return 0 < len(stripped) <= 2 and all(ch in _BULLET_CHARS for ch in stripped)


def _postprocess_bullet_tables(output_docx):
    """pdf2docx has a known quirk: a bulleted list in the source PDF (marker
    glyph + text sitting in two visually-separate columns) frequently gets
    reconstructed as a borderless 1-row x 2-column *table* — one cell full of
    stacked bullet glyphs, the other full of stacked item text — rather than
    real Word bulleted paragraphs. That's structurally wrong (screen readers,
    "List Bullet" styling, and further editing all break) even though it
    looks fine at a glance. This rewrites any table matching that exact shape
    into genuine bulleted paragraphs, in place, preserving their position in
    the document. Best-effort: any table that doesn't match is left alone,
    and any error here silently leaves the original pdf2docx output as-is
    rather than failing the whole conversion over a cosmetic pass.
    """
    try:
        from docx import Document

        doc = Document(output_docx)
        changed = False

        for table in list(doc.tables):
            try:
                if len(table.rows) != 1 or len(table.columns) != 2:
                    continue
                row = table.rows[0]
                marker_lines = row.cells[0].text.split("\n")
                item_lines = row.cells[1].text.split("\n")
                if len(marker_lines) < 1 or len(marker_lines) != len(item_lines):
                    continue
                if not all(_looks_like_bullet_marker(m) for m in marker_lines):
                    continue
                if not any(item.strip() for item in item_lines):
                    continue

                tbl_el = table._tbl
                for item in item_lines:
                    item = item.strip()
                    if not item:
                        continue
                    new_p = doc.add_paragraph(item, style="List Bullet")
                    # add_paragraph() appends at the end of the body; move it
                    # back to where the table used to be so document order
                    # (and surrounding context) is preserved.
                    tbl_el.addprevious(new_p._p)
                tbl_el.getparent().remove(tbl_el)
                changed = True
            except Exception:
                continue  # leave this particular table untouched

        if changed:
            doc.save(output_docx)
    except Exception as exc:
        # Never let the cleanup pass turn a successful conversion into a
        # failure — the un-postprocessed pdf2docx output is still valid.
        print(f"bullet-table postprocess skipped: {exc}", file=sys.stderr)


def main():
    if len(sys.argv) != 3:
        print("usage: pdf_to_docx.py <input.pdf> <output.docx>", file=sys.stderr)
        return 1

    input_pdf, output_docx = sys.argv[1], sys.argv[2]

    # pdf2docx logs one INFO line per stage ("Parsing pages...", etc). That's
    # noise for a CLI wrapper whose only real signal is the exit code, so mute
    # it down to warnings/errors.
    logging.getLogger("pdf2docx").setLevel(logging.WARNING)
    logging.getLogger("fitz").setLevel(logging.WARNING)

    try:
        from pdf2docx import Converter
    except ImportError as exc:
        print(f"pdf2docx is not installed: {exc}", file=sys.stderr)
        return 1

    try:
        cv = Converter(input_pdf)
    except Exception as exc:
        print(f"failed to open PDF: {exc}", file=sys.stderr)
        return 1

    try:
        # multi_processing=False: one page at a time. This runs inside a
        # per-request temp dir under the Node server's own concurrency
        # semaphore (MAX_CONCURRENT_JOBS), so we don't want pdf2docx forking
        # its own worker pool on top of that.
        cv.convert(output_docx, start=0, end=None, multi_processing=False)
    except Exception as exc:
        print(f"conversion failed: {exc}", file=sys.stderr)
        return 1
    finally:
        try:
            cv.close()
        except Exception:
            pass

    _postprocess_bullet_tables(output_docx)

    return 0


if __name__ == "__main__":
    sys.exit(main())
